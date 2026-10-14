"""QC report formatting, shared with the earlier ADNOC-style adaptation."""
from openpyxl.styles import Font, PatternFill, Alignment

def format_sheet(ws, header=1):
    ws.freeze_panes = "B" + str(header + 1)
    ws.auto_filter.ref = "A" + str(header) + ":" + ws.cell(ws.max_row, ws.max_column).coordinate
    ws.row_dimensions[header].height = 34
    for cell in ws[header]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=header + 1):
        ws.row_dimensions[row[0].row].height = 30
        for cell in row:
            if isinstance(cell.value, str):
                # Values from a GDB are text, never executable Excel formulas.
                cell.value = cell.value[:32767]
                cell.data_type = "s"
            if not cell.hyperlink:
                cell.font = Font(name="Calibri", size=11, color="808080" if cell.value == 0 else "000000", bold=isinstance(cell.value, (int, float)) and cell.value != 0)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column in ws.columns:
        width = min(48, max(16, max(len(str(c.value or "")) for c in column) + 3))
        ws.column_dimensions[column[0].column_letter].width = width
    for row in ws.iter_rows(min_row=header + 1):
        max_lines = 1
        for cell in row:
            width = ws.column_dimensions[cell.column_letter].width
            max_lines = max(max_lines, (len(str(cell.value or "")) // max(1, int(width) - 2)) + 1)
        ws.row_dimensions[row[0].row].height = min(90, max(30, max_lines * 15))
