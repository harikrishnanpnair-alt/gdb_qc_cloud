"""First open-source cloud GDB QC engine. Every source is opened read-only."""
from __future__ import annotations
import hashlib
import json
import math
import os
import re
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from osgeo import gdal, ogr
import shapely
from shapely.ops import transform as transform_geometry
from pyproj import CRS, Transformer
from cloud_gdb_common import open_gdb, inventory, resolve_schema, layer_definition, layer_fields
from report_style import format_sheet

MAX_FEATURES = int(os.environ.get("QC_MAX_TOTAL_FEATURES", "250000"))
DETAIL_LIMIT = int(os.environ.get("QC_MAX_DETAIL_ROWS", "50000"))


def code_key(value, domain):
    kind = domain.GetFieldType()
    if kind in {ogr.OFTInteger, ogr.OFTInteger64}:
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value != int(value):
            return None
        return str(int(value))
    if kind == ogr.OFTString:
        return value if isinstance(value, str) else None
    if kind == ogr.OFTReal:
        try:
            number = float(value)
            return str(number) if math.isfinite(number) else None
        except (ValueError, TypeError):
            return None
    raise ValueError("Domain field type is not supported by this first cloud version.")


def domain_violation(value, domain):
    if value is None:
        return None
    if domain.GetDomainType() == ogr.OFDT_CODED:
        codes = domain.GetEnumeration()
        key = code_key(value, domain)
        if domain.GetFieldType() == ogr.OFTReal:
            try:
                if key is not None and float(key) in {float(c) for c in codes}:
                    return None
            except (ValueError, TypeError):
                pass
        elif key is not None and key in codes:
            return None
        hint = next((c for c, desc in codes.items() if desc is not None and str(value).strip().lower() == str(desc).strip().lower()), None)
        return "Description used instead of code " + hint if hint is not None else "Invalid coded value"
    if domain.GetDomainType() == ogr.OFDT_RANGE:
        if domain.GetFieldType() not in {ogr.OFTInteger, ogr.OFTInteger64, ogr.OFTReal}:
            raise ValueError("Non-numeric range domains are not supported in this version.")
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            return "Value type does not match numeric range domain"
        low, high = domain.GetMinAsDouble(), domain.GetMaxAsDouble()
        if value < low or value > high or (value == low and not domain.IsMinInclusive()) or (value == high and not domain.IsMaxInclusive()):
            return "Outside range " + str(low) + " to " + str(high)
        return None
    raise ValueError("Glob/other domain kinds are not supported in this version.")


def length_calculator(layer):
    sr = layer.GetSpatialRef()
    if sr is None:
        raise ValueError("Short-line check needs a defined coordinate system.")
    crs = CRS.from_wkt(sr.ExportToWkt())
    geographic = crs.geodetic_crs
    if geographic is None:
        raise ValueError("Short-line check needs a valid geographic base CRS.")
    inverse = Transformer.from_crs(crs, geographic, always_xy=True)
    geod = geographic.get_geod()
    def length(geometry):
        return abs(geod.geometry_length(transform_geometry(inverse.transform, shapely.force_2d(geometry))))
    return length


def compare_schema(source_layer, target_layer, source_meta, target_meta):
    actual = {f["name"].lower(): f for f in layer_fields(source_layer, source_meta)}
    expected = layer_fields(target_layer, target_meta)
    issues = []
    for field in expected:
        found = actual.get(field["name"].lower())
        if found is None:
            issues.append([field["name"], "Missing field", "", field["type"]])
            continue
        for key, label in [("type", "Field type mismatch"), ("length", "Field length mismatch"),
                           ("nullable", "Nullable setting mismatch"), ("domain", "Domain assignment mismatch")]:
            if (key != "length" or field["ogr_type"] == ogr.OFTString) and found[key] != field[key]:
                issues.append([field["name"], label, str(found[key]), str(field[key])])
    if source_layer.GetGeomType() != target_layer.GetGeomType():
        issues.append(["Shape", "Geometry/Z/M type mismatch", ogr.GeometryTypeToName(source_layer.GetGeomType()), ogr.GeometryTypeToName(target_layer.GetGeomType())])
    a, b = source_layer.GetSpatialRef(), target_layer.GetSpatialRef()
    if a is None or b is None or not a.IsSame(b):
        issues.append(["Spatial Reference", "CRS mismatch", a.GetName() if a else "Unknown", b.GetName() if b else "Unknown"])
    if source_meta is not None and target_meta is not None:
        sf = target_meta["subtype_field"]
        if source_meta["subtype_field"].lower() != sf.lower():
            issues.append([sf, "Subtype field mismatch", source_meta["subtype_field"], sf])
        if sf:
            if source_meta["default_subtype"] != target_meta["default_subtype"]:
                issues.append([sf, "Default subtype mismatch", source_meta["default_subtype"], target_meta["default_subtype"]])
            for code, sub in target_meta["subtypes"].items():
                original = source_meta["subtypes"].get(code)
                if original is None:
                    issues.append([sf, "Missing subtype code", "", code])
                else:
                    for name, domain in sub["domains"].items():
                        actual_domain = original["domains"].get(name, "")
                        if actual_domain != domain:
                            issues.append([name, "Subtype domain assignment mismatch (" + code + ")", actual_domain, domain])
    return issues, expected


def geometry_from_ogr(geometry):
    if geometry.HasCurveGeometry():
        raise ValueError("Curved geometry: short-line and duplicate checks are not evaluated by this version.")
    result = shapely.from_wkb(bytes(geometry.ExportToIsoWkb()))
    if ogr.GT_HasM(geometry.GetGeometryType()) and not shapely.has_m(result):
        raise ValueError("Measured geometry could not retain M values for duplicate comparison.")
    return result


def run_qc(input_gdb, schema_gdb, output_excel, short_limit=2.0, selected=None, progress=None):
    input_items, schema_items = inventory(input_gdb), inventory(schema_gdb)
    if selected is not None:
        if not selected or any(p not in {r["path"] for r in input_items} for p in selected):
            raise ValueError("Select at least one inspected feature class.")
        input_items = [r for r in input_items if r["path"] in selected]
    if not input_items:
        raise ValueError("The input GDB has no spatial feature classes.")
    count = sum(max(0, r["count"]) for r in input_items)
    if count > MAX_FEATURES:
        raise ValueError("Selected feature count exceeds " + str(MAX_FEATURES) + ". Select fewer classes for this cloud trial.")
    source_ds, schema_ds = open_gdb(input_gdb), open_gdb(schema_gdb)
    domains = {name: schema_ds.GetFieldDomain(name) for name in (schema_ds.GetFieldDomainNames() or [])}
    wb = Workbook()
    executive = wb.active; executive.title = "Executive_Summary"
    executive.append(["Feature Class", "Status", "Known Issue Indicators", "Incomplete Checks", "Records"])
    summary = wb.create_sheet("Summary")
    headers = ["Feature Class", "Geometry", "Records", "Fields With Nulls/Blanks", "Null Cells", "Blank Cells", "Short Lines", "Exact Duplicate Extras", "Empty Geometry", "Domain Violations", "Schema Issues", "Status", "Acquisition_Method Values", "Data_Source Values", "Open Detail"]
    summary.append(headers)
    coverage = wb.create_sheet("Check_Coverage"); coverage.append(["Feature Class", "Check", "Reason"])
    notes = wb.create_sheet("Run_Details"); notes.append(["Setting", "Value"])
    notes.append(["Engine", "GDAL " + gdal.VersionInfo("RELEASE_NAME") + "; OGR/OpenFileGDB; Shapely " + shapely.__version__])
    notes.append(["Short-line threshold (m)", short_limit])
    notes.append(["Short-line length", "Geodesic XY length on the input CRS ellipsoid, using its inverse projection. No datum transformation to WGS84 is applied."])
    notes.append(["Duplicate rule", "Exact normalized geometry WKB, preserving XY/Z/M when supported. Direction/part order normalized; first encountered FID is the original. ArcGIS XY/Z snapping tolerance is NOT applied; results may differ from FindIdentical."])
    notes.append(["Issue indicators", "Null-field counts are indicators; null/blank cells are separate measures. Totals are not counts of distinct affected features."])
    notes.append(["Details limit", "At most " + str(DETAIL_LIMIT) + " detail rows for the workbook; summary counts include all scanned records. Truncation is listed on Check_Coverage."])
    notes.append(["Outside scope", "Topology rules, network connectivity, relationship/attachment integrity, nonspatial tables and project-specific rules outside the reference schema are not checked."])
    results = []
    details_used = 0
    scanned = 0
    for index, item in enumerate(input_items):
        layer = source_ds.GetLayerByName(item["name"])
        errors = {}
        details = {"Nulls": [], "Short": [], "Duplicates": [], "Empty": [], "Schema": [], "Domains": []}
        truncations = Counter()
        def incomplete(check, reason):
            errors.setdefault(check, str(reason))
        def detail(kind, row):
            nonlocal details_used
            if details_used < DETAIL_LIMIT:
                details[kind].append(row); details_used += 1
            else:
                truncations[kind] += 1
        try:
            source_meta = layer_definition(source_ds, layer)
        except Exception as exc:
            source_meta = None; incomplete("Subtype/schema metadata", exc)
        input_fields = layer_fields(layer, source_meta)
        field_map = {f["name"].lower(): f for f in input_fields}
        nulls = [[0, 0] for _ in input_fields]
        values_sets = {"acquisition_method": set(), "data_source": set()}
        expected_fields = []
        target_meta = None
        schema_issues = None
        domain_count = None
        try:
            target = resolve_schema(item, schema_items)
            if target is None:
                raise ValueError("Feature class is missing from the reference schema GDB.")
            target_layer = schema_ds.GetLayerByName(target["name"])
            try:
                target_meta = layer_definition(schema_ds, target_layer)
            except Exception as exc:
                incomplete("Subtype/schema metadata", exc)
            issues, expected_fields = compare_schema(layer, target_layer, source_meta, target_meta)
            schema_issues = len(issues)
            for row in issues:
                detail("Schema", row)
            if target_meta is None:
                incomplete("Domain values", "Reference subtype metadata could not be read, so all domain assignments could not be verified.")
            elif target_meta["subtype_field"] and target_meta["subtype_field"].lower() not in field_map:
                incomplete("Domain values", "The reference subtype field is missing from the input class.")
            needed_domains = {f["domain"] for f in expected_fields if f["domain"]}
            if target_meta:
                needed_domains.update(d for sub in target_meta["subtypes"].values() for d in sub["domains"].values() if d)
            for name in needed_domains:
                if domains.get(name) is None:
                    incomplete("Domain values", "Reference domain not available: " + name)
            domain_count = 0
        except Exception as exc:
            incomplete("Schema/domains", exc)
        empty_count = short_count = duplicate_count = 0
        seen_geometry = {}
        native_geom_type = ogr.GT_Flatten(layer.GetGeomType())
        is_line = native_geom_type in {ogr.wkbLineString, ogr.wkbMultiLineString, ogr.wkbCircularString, ogr.wkbCompoundCurve, ogr.wkbMultiCurve}
        length = None
        if is_line:
            try:
                length = length_calculator(layer)
            except Exception as exc:
                incomplete("Short lines", exc)
        layer.ResetReading()
        records = 0
        for feature in layer:
            records += 1; scanned += 1
            if scanned > MAX_FEATURES:
                raise ValueError("Actual scanned feature count exceeds the configured limit.")
            oid = feature.GetFID()
            for j, field in enumerate(input_fields):
                value = feature.GetField(field["index"])
                if value is None:
                    nulls[j][0] += 1
                elif isinstance(value, str) and not value.strip():
                    nulls[j][1] += 1
                key = field["name"].lower()
                if key in values_sets and len(values_sets[key]) < 200:
                    values_sets[key].add("<NULL>" if value is None or (isinstance(value, str) and not value.strip()) else str(value))
            if domain_count is not None:
                subtype = None
                if target_meta and target_meta["subtype_field"]:
                    subtype_field = field_map.get(target_meta["subtype_field"].lower())
                    if subtype_field:
                        subtype_value = feature.GetField(subtype_field["index"])
                        subtype = target_meta["subtypes"].get(str(subtype_value))
                        if subtype is None:
                            domain_count += 1
                            detail("Domains", [oid, subtype_field["name"], subtype_value, "Subtype", "Invalid or missing subtype code"])
                for expected in expected_fields:
                    actual = field_map.get(expected["name"].lower())
                    if actual is None:
                        continue
                    domain_name = expected["domain"]
                    if subtype and expected["name"].lower() in subtype["domains"]:
                        domain_name = subtype["domains"][expected["name"].lower()]
                    if not domain_name or domains.get(domain_name) is None:
                        continue
                    try:
                        value = feature.GetField(actual["index"])
                        reason = domain_violation(value, domains[domain_name])
                        if reason:
                            domain_count += 1
                            detail("Domains", [oid, actual["name"], value, domain_name, reason])
                    except Exception as exc:
                        incomplete("Domain values", str(exc) + " Domain: " + domain_name)
            geometry = feature.GetGeometryRef()
            if geometry is None or geometry.IsEmpty():
                empty_count += 1; detail("Empty", [oid]); continue
            try:
                shape = geometry_from_ogr(geometry)
            except Exception as exc:
                incomplete("Exact duplicates", exc)
                if is_line:
                    incomplete("Short lines", exc)
                continue
            if is_line and length is not None:
                try:
                    metres = length(shape)
                    if not math.isfinite(metres):
                        raise ValueError("Geodesic length is not finite.")
                    if metres < short_limit:
                        short_count += 1; detail("Short", [oid, metres])
                except Exception as exc:
                    incomplete("Short lines", exc)
            try:
                canonical = shapely.to_wkb(shapely.normalize(shape), byte_order=1, output_dimension=4)
                digest = hashlib.sha256(canonical).digest()
                if digest in seen_geometry:
                    duplicate_count += 1; detail("Duplicates", [oid, seen_geometry[digest]])
                else:
                    seen_geometry[digest] = oid
            except Exception as exc:
                incomplete("Exact duplicates", exc)
        null_rows = [[f["name"], f["alias"], f["type"], counts[0], counts[1]] for f, counts in zip(input_fields, nulls) if any(counts)]
        for row in null_rows:
            detail("Nulls", row)
        metrics = {"null_fields": len(null_rows), "null_cells": sum(r[3] for r in null_rows), "blank_cells": sum(r[4] for r in null_rows),
                   "short_segments": None if "Short lines" in errors else short_count,
                   "duplicates": None if "Exact duplicates" in errors else duplicate_count,
                   "empty_geometry": empty_count,
                   "domain_violations": None if "Domain values" in errors or "Schema/domains" in errors else domain_count,
                   "schema_issues": schema_issues}
        total = sum(metrics[k] or 0 for k in ("null_fields", "short_segments", "duplicates", "empty_geometry", "domain_violations", "schema_issues"))
        status = "INCOMPLETE" if errors else "Critical" if total >= 15 else "Warning" if total >= 5 else "Minor" if total else "No issues found"
        for check, reason in errors.items():
            coverage.append([item["path"], check, reason])
        for kind, omitted in truncations.items():
            coverage.append([item["path"], kind + " detail output", str(omitted) + " detail rows omitted; summary counts are complete."])
        definitions = {"Nulls": ["Field", "Alias", "Type", "Null Count", "Blank Count"], "Short": ["FID", "Length (m)"],
                       "Duplicates": ["Duplicate FID", "Original FID"], "Empty": ["FID"],
                       "Schema": ["Field", "Issue", "Actual", "Expected"], "Domains": ["FID", "Field", "Value", "Domain", "Issue"]}
        sheet_names = []
        for kind, rows in details.items():
            if not rows:
                continue
            name = (str(index+1) + "_" + re.sub(r"[\\/*?:\[\]]", "_", item["name"])[:16] + "_" + kind)[:31]
            ws = wb.create_sheet(name); ws.append(["Back to Summary"]); ws["A1"].hyperlink = "#'Summary'!A1"
            ws["A1"].font = Font(color="0000FF", underline="single"); ws.append(definitions[kind])
            for row in rows:
                ws.append([str(v) if isinstance(v, (list, dict, tuple)) else v for v in row])
            format_sheet(ws, 2); sheet_names.append(name)
        value_strings = [", ".join(sorted(values_sets[n])) + (" [at most 200 values shown]" if len(values_sets[n]) >= 200 else "") if n in field_map else "<FIELD MISSING>" for n in ("acquisition_method", "data_source")]
        summary.append([item["path"], item["geometry"], records] + [metrics[k] if metrics[k] is not None else "NOT CHECKED" for k in ("null_fields", "null_cells", "blank_cells", "short_segments", "duplicates", "empty_geometry", "domain_violations", "schema_issues")] + [status] + value_strings + ["Open" if sheet_names else ""])
        if sheet_names:
            cell = summary.cell(summary.max_row, 15); cell.hyperlink = "#'" + sheet_names[0] + "'!A1"; cell.font = Font(color="0000FF", underline="single")
        executive.append([item["path"], status, total, len(errors), records])
        results.append({"path": item["path"], "geometry": item["geometry"], "records": records, **metrics,
                        "status": status, "known_issue_indicators": total,
                        "check_failures": [{"check": k, "error": v} for k, v in errors.items()],
                        "omitted_detail_rows": dict(truncations)})
        if progress:
            progress((index+1)/len(input_items), "Checked " + item["path"])
    for ws in (executive, summary, coverage, notes):
        format_sheet(ws)
    for ws, column in [(executive, 2), (summary, 12)]:
        for row in ws.iter_rows(min_row=2):
            status = row[column-1].value
            if status != "No issues found":
                row[column-1].fill = PatternFill("solid", fgColor={"INCOMPLETE":"F4B183","Critical":"FF9999","Warning":"FFF2CC","Minor":"E2F0D9"}[status])
    wb.save(output_excel)
    source_ds = schema_ds = None
    return {"engine": "GDAL/OGR", "classes": results, "status_counts": dict(Counter(r["status"] for r in results)),
            "has_warnings": any(r["status"] == "INCOMPLETE" or r["omitted_detail_rows"] for r in results),
            "short_limit_m": short_limit, "duplicate_rule": "Exact normalized geometry, without ArcGIS snapping tolerance"}
